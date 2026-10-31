"""Actual PDF/DOCX text extraction; optional per-page PDF OCR using local Tesseract."""
import io
from dataclasses import dataclass
from pathlib import Path
import pdfplumber
from docx import Document
from app.core.config import get_settings
from app.documents.validation import DOCX_MIME, PDF_MIME

PARSER_VERSION = 'phase3-parser-v1'


@dataclass
class PageResult:
    number: int
    text: str
    method: str
    warnings: list[str]


def extract_pdf(path: str) -> list[PageResult]:
    settings = get_settings()
    pages: list[PageResult] = []
    with pdfplumber.open(path) as doc:
        if len(doc.pages) > settings.max_pdf_pages:
            raise ValueError('PDF_PAGE_LIMIT')
        for i, p in enumerate(doc.pages, 1):
            text = (p.extract_text() or '').strip()
            warnings = []
            method = 'pdfplumber'
            if len(text) < 35:
                # Only OCR pages without usable text; no invented content if OCR fails.
                try:
                    import pypdfium2 as pdfium
                    import pytesseract
                    pdf = pdfium.PdfDocument(path)
                    rendered = pdf[i-1].render(scale=1.7).to_pil()
                    scanned = pytesseract.image_to_string(rendered, lang=settings.ocr_languages, timeout=25).strip()
                    if len(scanned) > len(text):
                        text = scanned
                        method = 'tesseract'
                    else:
                        warnings.append('OCR_EMPTY')
                except Exception:
                    warnings.append('OCR_UNAVAILABLE')
            if not text:
                warnings.append('NO_TEXT')
            pages.append(PageResult(i, text, method, warnings))
    return pages


def extract_docx(path: str) -> list[PageResult]:
    settings = get_settings()
    document = Document(path)
    paragraphs = list(document.paragraphs)
    if len(paragraphs) > settings.max_docx_paragraphs:
        raise ValueError('DOCX_PARAGRAPH_LIMIT')
    text_parts = [p.text for p in paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            text_parts.append(' | '.join(cell.text for cell in row.cells))
    text = '\n'.join(text_parts).strip()
    # DOCX does not expose stable pagination; page 1 denotes logical document, not physical page.
    return [PageResult(1, text, 'python-docx', [] if text else ['NO_TEXT'])]


def extract_document(path: str, mime: str) -> list[PageResult]:
    if mime == PDF_MIME:
        return extract_pdf(path)
    if mime == DOCX_MIME:
        return extract_docx(path)
    raise ValueError('UNSUPPORTED_FORMAT')
